#!/usr/bin/env python3
# ©  2026 Terrabit Solutions SRL
#    Dorin Hongu <dhongu(@)gmail(.)com>
"""Generează nomenclatorul de coduri UNECE din schema SAF-T publicată de ANAF.

Sursa e fila ``Unitati_masura`` din schema SAF-T publicată de ANAF:
https://static.anaf.ro/static/10/Anaf/Informatii_R/RO_SAFT_SchemaDefCod_16.02.2026.xlsx
Fila listează codurile admise la raportare, cu denumirea în engleză și traducerea
indicativă în română. Fișierul nu e versionat în repo — are 2 MB și se republică
periodic — deci scriptul primește calea lui ca argument.

Scrie ``data/uom.unece.code.csv``: codul, sursa (rec20/rec21), denumirea engleză
și, în coloana ``name@ro``, pe cea română. Sintaxa ``<câmp>@<limbă>`` e cea din
``l10n_ro/data/template/account.account-ro.csv`` — Odoo scoate coloana din datele
propriu-zise la import și o reia ca traducere, legată de xml_id.

Denumirile RO NU mai trec prin ``i18n/ro.po``: acolo traducerea se potrivește
prin msgid via ``.pot``, iar un ``.pot`` rămas în urmă o lipește tăcut pe altă
înregistrare. ``ro.po`` rămâne doar pentru etichetele de interfață.

    python3 scripts/import_unece_from_saft.py ~/Downloads/RO_SAFT_SchemaDefCod_16.02.2026.xlsx
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("openpyxl lipsește: pip install openpyxl")

SHEET = "Unitati_masura"
FIRST_DATA_ROW = 12  # rândurile 1-11 sunt copertă, date de publicare și antet

# Corecții față de fișierul ANAF, fiecare cu motivul ei. Nomenclatorul e o sursă
# oficială, dar nu e lipsit de erori, iar noi nu putem propaga mai departe un
# text despre care știm că e greșit.
NAME_RO_FIXES = {
    # „square foot" e tradus în fișier tot prin „metru pătrat", ca FTK să nu iasă
    # identic cu MTK în interfață.
    "FTK": "picior pătrat",
    # Rămase netraduse în fișier.
    "C62": "unu (bucată)",
    # Tradus automat ca „negator": „denier" e unitatea de finețe a firelor
    # textile, nu substantivul englezesc.
    "A49": "denier",
    "M83": "denier",
}

# Coduri care apar de două ori în fișier. Păstrăm varianta marcată ca modificare
# față de publicarea anterioară, deci ultima apariție din filă.
KNOWN_DUPLICATES = {"B30"}


def read_rows(xlsx: Path):
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        sys.exit(f"fila {SHEET!r} lipsește din {xlsx.name}")
    seen, rows = {}, []
    for index, row in enumerate(wb[SHEET].iter_rows(values_only=True), 1):
        if index < FIRST_DATA_ROW or not row[1]:
            continue
        code = str(row[1]).strip()
        source = str(row[0] or "").strip() or "other"
        name_en = (row[2] or "").strip()
        name_ro = (row[3] or "").strip() or name_en
        if not name_en:
            print(f"  ! {code}: fără denumire engleză, sărit")
            continue
        entry = (code, source, name_en, NAME_RO_FIXES.get(code, name_ro))
        if code in seen:
            if code not in KNOWN_DUPLICATES:
                print(f"  ! {code}: duplicat nedocumentat, păstrez ultima apariție")
            rows[seen[code]] = entry
        else:
            seen[code] = len(rows)
            rows.append(entry)
    return rows


def write_csv(rows, path: Path):
    """Scrie CSV-ul de date, cu denumirea engleză și cea română pe același rând.

    Coloana ``name@ro`` e sintaxa Odoo pentru traducerea unui câmp direct în
    fișierul de date — aceeași folosită de ``l10n_ro/data/template/
    account.account-ro.csv`` pentru planul de conturi. La instalare,
    ``convert_csv_import`` scoate coloanele cu ``@`` din datele propriu-zise, iar
    ``CSVDataFileReader`` le reia ca traduceri.

    Avantajul față de a ține denumirile RO în ``i18n/ro.po``: traducerea se leagă
    de **xml_id**, nu prin potrivire de msgid prin ``.pot``. Un ``.pot`` rămas în
    urmă nu mai poate lipi traducerea pe altă înregistrare — exact ce s-a
    întâmplat înainte, când XPF („Pen") a ieșit numit „Palet".
    """
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(["id", "code", "name", "source", "name@ro"])
        for code, source, name_en, name_ro in rows:
            writer.writerow([f"unece_{code.lower()}", code, name_en, source, name_ro])
    print(f"CSV   → {path}  ({len(rows)} coduri, EN + RO)")


def po_escape(value: str) -> str:
    return value.replace("\\", "\\\\").replace('"', '\\"').replace("\n", "\\n")


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("xlsx", type=Path, help="Schema SAF-T publicată de ANAF (.xlsx)")
    parser.add_argument(
        "--addon-dir",
        type=Path,
        default=Path(__file__).resolve().parent.parent / "deltatech_uom_unece",
        help="Modulul în care se scriu fișierele",
    )
    args = parser.parse_args()
    if not args.xlsx.exists():
        sys.exit(f"nu găsesc {args.xlsx}")

    rows = read_rows(args.xlsx)
    odd = [code for code, *_ in rows if not re.fullmatch(r"[0-9A-Z]{2,4}", code)]
    if odd:
        sys.exit(f"coduri cu format neașteptat: {odd}")
    long_codes = [code for code, *_ in rows if len(code) > 3]
    if long_codes:
        # Schema eTransport (`CodUMType`) acceptă doar 2-3 caractere; le raportăm
        # ca să se știe că nu orice cod din SAF-T e trimisibil pe orice document.
        print(f"  i coduri de 4 caractere, neutilizabile pe eTransport: {long_codes}")

    write_csv(rows, args.addon_dir / "data" / "uom.unece.code.csv")


if __name__ == "__main__":
    main()
