# ©  2026 Deltatech
#              Dorin Hongu <dhongu(@)gmail(.)com
# See README.rst file on addons root folder for license details

import base64
import io
from unittest.mock import patch

import openpyxl

from odoo.exceptions import UserError
from odoo.tests import TransactionCase, tagged


@tagged("post_install", "-at_install")
class TestPurchaseXls(TransactionCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.env = cls.env(context=dict(cls.env.context, tracking_disable=True))
        cls.vendor = cls.env["res.partner"].create({"name": "PXLS Vendor", "is_company": True})
        cls.other_vendor = cls.env["res.partner"].create({"name": "PXLS Other Vendor", "is_company": True})
        cls.uom_unit = cls.env.ref("uom.product_uom_unit")
        cls.uom_kg = cls.env.ref("uom.product_uom_kgm")
        cls.product_a = cls.env["product.product"].create(
            {"name": "PXLS Product A", "default_code": "PXA", "type": "consu", "is_storable": True}
        )
        cls.product_b = cls.env["product.product"].create(
            {"name": "PXLS Product B", "default_code": "PXB", "type": "consu", "is_storable": True}
        )
        cls.product_c = cls.env["product.product"].create({"name": "PXLS Product C", "type": "consu"})
        cls.env["product.supplierinfo"].create(
            {
                "partner_id": cls.vendor.id,
                "product_tmpl_id": cls.product_a.product_tmpl_id.id,
                "product_code": "SUP-A",
                "product_name": "Vendor name A",
                "price": 11.0,
            }
        )
        cls.env["product.supplierinfo"].create(
            {
                "partner_id": cls.other_vendor.id,
                "product_tmpl_id": cls.product_b.product_tmpl_id.id,
                "product_code": "OTHER-B",
                "price": 12.0,
            }
        )

    # helpers

    def _create_order(self, products=None, partner=None):
        lines = [(0, 0, {"product_id": p.id, "product_qty": 1.0, "price_unit": 5.0}) for p in products or []]
        return self.env["purchase.order"].create({"partner_id": (partner or self.vendor).id, "order_line": lines})

    @staticmethod
    def _xlsx(rows):
        book = openpyxl.Workbook()
        sheet = book.active
        for row in rows:
            sheet.append(row)
        buffer = io.BytesIO()
        book.save(buffer)
        return base64.b64encode(buffer.getvalue())

    def _wizard(self, order, rows, **values):
        values.setdefault("fields_list", "product_code,product_name,quantity,price,uom_name")
        return (
            self.env["import.purchase.line"]
            .with_context(active_id=order.id, active_model="purchase.order")
            .create(dict(values, data_file=self._xlsx(rows)))
        )

    # purchase.order

    def test_show_order_lines(self):
        order = self._create_order([self.product_a])
        action = order.show_order_lines()
        self.assertEqual(action["type"], "ir.actions.act_window")
        self.assertEqual(action["res_model"], "purchase.order.line")
        self.assertEqual(action["domain"], [("order_id", "=", order.id)])
        self.assertEqual(action["context"]["default_order_id"], order.id)
        view = self.env.ref("deltatech_purchase_xls.purchase_order_line_tree")
        self.assertEqual(action["views"], [(view.id, "list")])

    # purchase.order.line load

    def test_load_update_existing_lines(self):
        """Lines matching products of the order are updated, the others are dropped."""
        order = self._create_order([self.product_a, self.product_c])
        line_a = order.order_line.filtered(lambda line: line.product_id == self.product_a)
        line_c = order.order_line.filtered(lambda line: line.product_id == self.product_c)
        fields = ["product_id", "product_qty"]
        data = [
            ["[PXA] PXLS Product A", "7"],  # matched by code
            ["PXLS Product C", "4"],  # matched by name
            ["[PXB] PXLS Product B", "3"],  # product exists but not on order
        ]
        result = self.env["purchase.order.line"].with_context(default_order_id=order.id).load(fields, data)
        self.assertFalse([m for m in result["messages"] if m["type"] == "error"], result["messages"])
        self.assertEqual(len(data), 2)
        self.assertIn(".id", fields)
        self.assertEqual(line_a.product_qty, 7.0)
        self.assertEqual(line_c.product_qty, 4.0)
        self.assertEqual(len(order.order_line), 2)

    def test_load_existing_lines_consecutive_dropped_rows(self):
        """Rows are removed from the list being iterated, so the row after a dropped one is skipped:
        it is neither matched nor dropped and keeps an empty .id (current behaviour)."""
        order = self._create_order([self.product_a])
        line = order.order_line
        fields = ["product_id", "product_qty"]
        data = [["[NOPE] Unknown", "1"], ["[PXB] PXLS Product B", "2"], ["[PXA] PXLS Product A", "3"]]
        with patch("odoo.models.BaseModel.load", autospec=True) as mock_load:
            mock_load.return_value = {"ids": [], "messages": []}
            self.env["purchase.order.line"].with_context(active_id=order.id).load(fields, data)
        mock_load.assert_called_once()
        self.assertEqual(data, [["[PXB] PXLS Product B", "2", ""], ["[PXA] PXLS Product A", "3", str(line.id)]])

    def test_load_existing_lines_without_product_column(self):
        order = self._create_order([self.product_a])
        fields = ["product_qty"]
        data = [["2"]]
        with patch("odoo.models.BaseModel.load", autospec=True) as mock_load:
            mock_load.return_value = {"ids": [], "messages": []}
            self.env["purchase.order.line"].with_context(active_id=order.id).load(fields, data)
        mock_load.assert_called_once()
        self.assertEqual(fields, ["product_qty", ".id"])
        self.assertEqual(data, [["2", ""]])

    def test_load_new_lines_in_empty_order(self):
        """In an order without lines the rows are created as new lines of the order."""
        order = self._create_order()
        fields = ["product_id", "product_qty", "price_unit"]
        data = [["PXLS Product A", "3", "9"], ["PXLS Product C", "2", "4"]]
        result = self.env["purchase.order.line"].with_context(default_order_id=order.id).load(fields, data)
        self.assertFalse([m for m in result["messages"] if m["type"] == "error"], result["messages"])
        self.assertEqual(fields[-1], "order_id")
        self.assertEqual(len(order.order_line), 2)
        line_a = order.order_line.filtered(lambda line: line.product_id == self.product_a)
        self.assertEqual(line_a.product_qty, 3.0)
        self.assertEqual(line_a.price_unit, 9.0)

    def test_load_without_order_context(self):
        """Without an order in context the code calls ``fields.index.get`` on a list (current behaviour)."""
        fields = ["product_id", "product_qty"]
        data = [["PXLS Product A", "1"]]
        with self.assertRaises(AttributeError):
            self.env["purchase.order.line"].load(fields, data)

    def test_parse_import_data(self):
        """The override only delegates to super, which BaseModel does not define."""
        with self.assertRaises(AttributeError):
            self.env["purchase.order.line"]._parse_import_data([["2"]], ["product_qty"], {})

    # import wizard

    def test_default_get_order_not_draft(self):
        order = self._create_order([self.product_a])
        order.button_confirm()
        with self.assertRaises(UserError):
            self._wizard(order, [["SUP-A", "x", 1, 1]])

    def test_get_rows_values(self):
        order = self._create_order()
        wizard = self._wizard(order, [["code", "name", "qty", "price"], ["SUP-A", None, 3, 2.5, True]], has_header=True)
        self.assertEqual(wizard.purchase_id, order)
        self.assertEqual(wizard.get_rows(), [["SUP-A", "", "3", "2.5", True]])

    def test_get_rows_empty_file(self):
        wizard = self._wizard(self._create_order(), [])
        self.assertFalse(wizard.get_rows())

    def test_get_rows_without_openpyxl(self):
        wizard = self._wizard(self._create_order(), [["SUP-A"]])
        with patch("odoo.addons.deltatech_purchase_xls.wizard.import_purchase_line.openpyxl", None):
            with self.assertRaises(UserError):
                wizard.get_rows()

    def test_import_supplier_code(self):
        order = self._create_order()
        rows = [
            ["code", "name", "qty", "price", "uom"],
            ["SUP-A", "Imported A", 4, 7.5, self.uom_unit.name],
            ["", "no code", 1, 1, ""],  # skipped: no product code
            ["SUP-A", "bad price", 1, "abc", ""],  # skipped: price is not a number
        ]
        self._wizard(order, rows, has_header=True).do_import()
        self.assertEqual(len(order.order_line), 1)
        line = order.order_line
        self.assertEqual(line.product_id, self.product_a)
        self.assertEqual(line.name, "Imported A")
        self.assertEqual(line.product_qty, 4.0)
        self.assertEqual(line.price_unit, 7.5)
        self.assertEqual(line.product_uom_id, self.uom_unit)

    def test_import_amount(self):
        order = self._create_order()
        self._wizard(order, [["SUP-A", "", 4, 10, ""]], is_amount=True).do_import()
        self.assertEqual(order.order_line.price_unit, 2.5)
        self.assertEqual(order.order_line.name, self.product_a.display_name)

    def test_import_unknown_uom_keeps_product_uom(self):
        order = self._create_order()
        self._wizard(order, [["SUP-A", "A", 1, 1, "PXLS no such uom"]]).do_import()
        self.assertEqual(order.order_line.product_uom_id, self.product_a.uom_id)

    def test_import_other_uom(self):
        order = self._create_order()
        wizard = self._wizard(order, [["SUP-A", "A", 1, 1, self.uom_kg.name]])
        with self.assertRaises(UserError):
            wizard.do_import()

    def test_import_product_not_found(self):
        order = self._create_order()
        wizard = self._wizard(order, [["OTHER-B", "B", 1, 1, ""]])
        with self.assertRaises(UserError):
            wizard.do_import()

    def test_import_search_by_default_code(self):
        order = self._create_order()
        wizard = self._wizard(
            order,
            [["PXB", "B", 2, 3]],
            fields_list="product_code,product_name,quantity,price",
            search_by_default_code=True,
        )
        self.assertFalse(wizard.search_product("PXLS-NOPE"))
        wizard.do_import()
        self.assertEqual(order.order_line.product_id, self.product_b)

    def test_import_without_price_column(self):
        """Without a price column the price is computed from the vendor pricelist."""
        order = self._create_order()
        self._wizard(order, [["SUP-A", 5]], fields_list="product_code,quantity").do_import()
        self.assertEqual(order.order_line.product_id, self.product_a)
        self.assertEqual(order.order_line.price_unit, 11.0)

    def test_import_invalid_fields_list(self):
        wizard = self._wizard(self._create_order(), [["SUP-A", "A", 1, 1]])
        wizard.fields_list = False
        with self.assertRaises(UserError):
            wizard.do_import()

    def test_import_new_product(self):
        order = self._create_order()
        rows = [["NEW-1", "PXLS New 1", 2, 6, self.uom_kg.name], ["NEW-2", "PXLS New 2", 1, 3, ""]]
        self._wizard(order, rows, new_product=True).do_import()
        self.assertEqual(len(order.order_line), 2)
        line_1 = order.order_line.filtered(lambda line: line.name == "PXLS New 1")
        self.assertEqual(line_1.product_id.uom_id, self.uom_kg)
        self.assertTrue(line_1.product_id.is_storable)
        seller = line_1.product_id.seller_ids
        self.assertEqual(seller.partner_id, self.vendor)
        self.assertEqual(seller.product_code, "NEW-1")
        self.assertEqual(seller.price, 6.0)
        self.assertEqual(seller.company_id, order.company_id)
        line_2 = order.order_line.filtered(lambda line: line.name == "PXLS New 2")
        self.assertEqual(line_2.product_id.uom_id.id, 1)

    # export wizard

    def test_export_vendor_code_and_name(self):
        order = self._create_order([self.product_a, self.product_b, self.product_c])
        wizard = (
            self.env["export.purchase.line"]
            .with_context(active_ids=order.ids, active_model="purchase.order")
            .create({})
        )
        action = wizard.do_export()
        self.assertEqual(action["res_id"], wizard.id)
        self.assertEqual(wizard.state, "get")
        book = openpyxl.load_workbook(io.BytesIO(base64.b64decode(wizard.data_file)))
        rows = list(book.worksheets[0].iter_rows(values_only=True))
        self.assertEqual(rows[0][:3], ("Code", "Internal Code", "Name"))
        by_internal = {row[1]: row for row in rows[1:]}
        self.assertEqual(by_internal["PXA"][0], "SUP-A")
        self.assertEqual(by_internal["PXA"][2], "Vendor name A")
        # seller of another vendor is ignored
        self.assertIn(by_internal["PXB"][0], ("", None))
        self.assertEqual(by_internal["PXB"][2], "PXLS Product B")
